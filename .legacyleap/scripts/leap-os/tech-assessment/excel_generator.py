#!/usr/bin/env python3
"""
Excel Generator: Technical Assessment Report
Combines all sheet JSONs into a single Excel workbook
"""

import json
import os
import re
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def _resolve_workspace():
    """Walk up from CWD to find .legacyleap/metadata.json and derive workspace root."""
    search = Path.cwd()
    while search != search.parent:
        meta_path = search / ".legacyleap" / "metadata.json"
        if meta_path.exists():
            return search, meta_path
        search = search.parent
    return Path.cwd(), None


WORKSPACE_ROOT, _META_FILE = _resolve_workspace()
_METADATA = {}
if _META_FILE and _META_FILE.exists():
    with open(_META_FILE) as _f:
        _METADATA = json.load(_f)

_scan_reports_raw = _METADATA.get("scan_assessment_reports", "scan_assessment_reports")
if os.path.isabs(_scan_reports_raw):
    _BASE_DIR = Path(_scan_reports_raw)
else:
    _BASE_DIR = WORKSPACE_ROOT / _scan_reports_raw

TECH_ASSESSMENT_DIR = _BASE_DIR / "tech_assessment"
OUTPUT_FILE = TECH_ASSESSMENT_DIR / "technical_assessment_report.xlsx"
SHEET1_FILE = TECH_ASSESSMENT_DIR / "sheet1.json"
SHEET2_FILE = TECH_ASSESSMENT_DIR / "sheet2.json"
SHEET3_FILE = TECH_ASSESSMENT_DIR / "sheet3.json"
SHEET4_FILE = TECH_ASSESSMENT_DIR / "sheet4.json"
SHEET7_FILE = TECH_ASSESSMENT_DIR / "sheet7.json"
SHEET8_FILE = TECH_ASSESSMENT_DIR / "sheet8.json"
SHEET9_FILE = TECH_ASSESSMENT_DIR / "sheet9.json"

# Styles
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)


def clean_text(text):
    """Remove control characters from text."""
    if not isinstance(text, str):
        return text
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)


def auto_adjust_column_width(sheet, min_width=10, max_width=50):
    """Auto-adjust column widths based on content."""
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        
        for cell in column_cells:
            try:
                cell_length = len(str(cell.value)) if cell.value else 0
                max_length = max(max_length, cell_length)
            except:
                pass
        
        adjusted_width = min(max(max_length + 2, min_width), max_width)
        sheet.column_dimensions[column_letter].width = adjusted_width


def apply_header_style(sheet, row=1):
    """Apply header styling to the first row."""
    for cell in sheet[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def merge_blocks(sheet, key_cols, first_data_row=2):
    """Merge cells in key columns when values are the same."""
    last_row = sheet.max_row
    if last_row < first_data_row:
        return
    
    def row_key(r):
        return tuple(sheet.cell(row=r, column=c).value for c in key_cols)
    
    def merge_range(col, start, end):
        if end > start:
            col_letter = get_column_letter(col)
            sheet.merge_cells(f"{col_letter}{start}:{col_letter}{end}")
            for r in range(start, end + 1):
                sheet[f"{col_letter}{r}"].alignment = Alignment(vertical="center")
    
    block_start = first_data_row
    current_key = row_key(first_data_row)
    
    for r in range(first_data_row + 1, last_row + 1):
        if row_key(r) != current_key:
            for col in key_cols:
                merge_range(col, block_start, r - 1)
            block_start = r
            current_key = row_key(r)
    
    for col in key_cols:
        merge_range(col, block_start, last_row)


def load_json_file(file_path):
    """Load JSON file if it exists."""
    path = Path(file_path)
    if not path.exists():
        print(f"  ⚠️ File not found: {file_path}")
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError) as exc:
        print(f"  ⚠️ Failed to read {file_path}: {exc}")
        return None


def create_sheet1(wb, data):
    """Create Sheet 1: Executive Decision Summary."""
    sheet = wb.create_sheet("1-Executive Summary")

    columns = data.get("columns", [
        "application_name", "current_version", "business_criticality",
        "technical_health_score", "modernization_readiness_score",
        "risk_level", "key_blocking_risks",
    ])

    readable_map = {
        "application_name": "Application Name",
        "current_version": "Current Version",
        "business_criticality": "Business Criticality",
        "technical_health_score": "Technical Health Score (0-100)",
        "modernization_readiness_score": "Modernization Readiness (0-100)",
        "risk_level": "Risk Level",
        "key_blocking_risks": "Key Blocking Risks",
    }
    readable_headers = [readable_map.get(col, col.replace("_", " ").title()) for col in columns]

    tech = data.get("technology", "")
    framework = data.get("framework", "")
    sheet.append([f"Executive Decision Summary  |  {tech} / {framework}"])
    sheet.merge_cells(f"A1:{get_column_letter(len(columns))}1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])

    sheet.append(readable_headers)
    apply_header_style(sheet, sheet.max_row)

    risk_fills = {
        "High": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
        "Medium": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
        "Critical": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
    }

    for row_data in data.get("data", []):
        row = [clean_text(row_data.get(col, "")) if isinstance(row_data.get(col, ""), str) else row_data.get(col, "") for col in columns]
        sheet.append(row)

        cur_row = sheet.max_row
        if "risk_level" in columns:
            rl_val = row_data.get("risk_level", "")
            if rl_val in risk_fills:
                rl_col = columns.index("risk_level") + 1
                sheet.cell(row=cur_row, column=rl_col).fill = risk_fills[rl_val]

        if "business_criticality" in columns:
            bc_val = row_data.get("business_criticality", "")
            if bc_val in risk_fills:
                bc_col = columns.index("business_criticality") + 1
                sheet.cell(row=cur_row, column=bc_col).fill = risk_fills[bc_val]

    auto_adjust_column_width(sheet)
    return sheet


def create_sheet2(wb, data):
    """Create Sheet 2: Modernization Decision Matrix."""
    sheet = wb.create_sheet("2-Modernization")

    columns = data.get("columns", [
        "module", "rehost_feasibility", "refactor_feasibility",
        "rearchitect_need", "replacement_indicator",
        "estimated_effort", "primary_drivers",
    ])

    readable_map = {
        "module": "Application / Module",
        "rehost_feasibility": "Re-host Feasibility",
        "refactor_feasibility": "Refactor Feasibility",
        "rearchitect_need": "Re-architect Need",
        "replacement_indicator": "Replacement Indicator",
        "estimated_effort": "Estimated Effort",
        "primary_drivers": "Primary Drivers",
    }
    readable_headers = [readable_map.get(col, col.replace("_", " ").title()) for col in columns]

    tech = data.get("technology", "")
    framework = data.get("framework", "")
    summary = data.get("summary", {})
    sheet.append([f"Modernization Decision Matrix  |  {tech} / {framework}  |  "
                  f"Modules: {summary.get('total_modules', 0)}"])
    sheet.merge_cells(f"A1:{get_column_letter(len(columns))}1")
    sheet["A1"].font = Font(bold=True, size=12)
    sheet.append([])

    sheet.append(readable_headers)
    apply_header_style(sheet, sheet.max_row)

    indicator_fills = {
        "Yes": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
        "Partial": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
    }

    for row_data in data.get("data", []):
        row = [clean_text(row_data.get(col, "")) if isinstance(row_data.get(col, ""), str) else row_data.get(col, "") for col in columns]
        sheet.append(row)

        if "replacement_indicator" in columns:
            ri_val = row_data.get("replacement_indicator", "")
            if ri_val in indicator_fills:
                ri_col = columns.index("replacement_indicator") + 1
                sheet.cell(row=sheet.max_row, column=ri_col).fill = indicator_fills[ri_val]

    auto_adjust_column_width(sheet)
    return sheet


def create_sheet3(wb, data):
    """Create Sheet 3: Application-Level Summary."""
    sheet = wb.create_sheet("3-App Summary")

    columns = data.get("columns", [
        "repo_name", "total_files", "legacy_files_pct", "high_risk_files_count",
        "architectural_violations", "security_issues", "test_coverage_indicator",
        "cloud_readiness_score",
    ])

    readable_map = {
        "repo_name": "Repo Name",
        "total_files": "Total Files",
        "legacy_files_pct": "Legacy Files (%)",
        "high_risk_files_count": "High-Risk Files Count",
        "architectural_violations": "Architectural Violations",
        "security_issues": "Security Issues",
        "test_coverage_indicator": "Test Coverage Indicator",
        "cloud_readiness_score": "Cloud Readiness Score",
    }
    readable_headers = [readable_map.get(col, col.replace("_", " ").title()) for col in columns]

    app_name = data.get("application_name", "")
    tech = data.get("technology", "")
    framework = data.get("framework", "")
    total_projects = data.get("total_projects_detected", 0)
    total_files = data.get("total_files_all_projects", 0)

    sheet.append([f"Application: {app_name}  |  Technology: {tech}  |  Framework: {framework}"])
    sheet.merge_cells(f"A1:{get_column_letter(len(columns))}1")
    sheet["A1"].font = Font(bold=True, size=12)

    sheet.append([f"Projects Detected: {total_projects}  |  Total Source Files: {total_files}"])
    sheet.merge_cells(f"A2:{get_column_letter(len(columns))}2")
    sheet["A2"].font = Font(italic=True)

    sheet.append([])

    sheet.append(readable_headers)
    apply_header_style(sheet, sheet.max_row)

    for row_data in data.get("data", []):
        row = [clean_text(row_data.get(col, "")) for col in columns]
        sheet.append(row)

    auto_adjust_column_width(sheet)
    return sheet


def create_sheet4(wb, data):
    """Create Sheet 4: Architecture Overview."""
    sheet = wb.create_sheet("4-Architecture")

    columns = data.get("columns", [
        "module_layer", "dependency_count", "cyclic_dependency",
        "cross_repo_dependency", "boundary_violation", "coupling_score",
        "architectural_risk", "change_impact",
    ])

    readable_map = {
        "module_layer": "Module / Layer",
        "dependency_count": "Dependency Count",
        "cyclic_dependency": "Cyclic Dependency",
        "cross_repo_dependency": "Cross-Repo Dependency",
        "boundary_violation": "Boundary Violation",
        "coupling_score": "Coupling Score",
        "architectural_risk": "Architectural Risk",
        "change_impact": "Change Impact",
    }
    readable_headers = [readable_map.get(col, col.replace("_", " ").title()) for col in columns]

    summary = data.get("summary", {})
    sheet.append([f"Architecture Overview  |  Modules: {summary.get('total_modules', 0)}  |  "
                  f"High Risk: {summary.get('high_risk', 0)}  |  "
                  f"Boundary Violations: {summary.get('boundary_violations', 0)}"])
    sheet.merge_cells(f"A1:{get_column_letter(len(columns))}1")
    sheet["A1"].font = Font(bold=True, size=12)
    sheet.append([])

    sheet.append(readable_headers)
    apply_header_style(sheet, sheet.max_row)

    risk_fills = {
        "High": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
        "Medium": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
    }

    for row_data in data.get("data", []):
        row = []
        for col in columns:
            val = row_data.get(col, "")
            row.append(clean_text(val) if isinstance(val, str) else val)
        sheet.append(row)

        risk = row_data.get("architectural_risk", "")
        if risk in risk_fills:
            risk_col = columns.index("architectural_risk") + 1
            sheet.cell(row=sheet.max_row, column=risk_col).fill = risk_fills[risk]

    auto_adjust_column_width(sheet)
    return sheet


def create_sheet7(wb, data):
    """Create Sheet 7: Code Quality & Maintainability."""
    sheet = wb.create_sheet("7-Code Quality")
    
    # Headers from the columns list
    columns = data.get("columns", [])
    if not columns:
        columns = [
            "filename", "class_name", "path", "vulnerabilities", "bugs", "code_smells",
            "coverage", "cognitive_complexity", "functions", "duplicate_lines_density",
            "lines", "security_risk_rating", "tech_debt_time", "architecture_complexity",
            "nfr_security_rating", "nfr_maintainability_rating", "nfr_reliability_rating",
            "nfr_testability_rating", "cloud_compatibiliy_rating", "coupling_score",
            "cloud_compatibiliy_assessment", "nfr_extensibility_rating", "nfr_reusability_rating",
            "nfr_scalability_rating", "risk_category", "mcm_node_id"
        ]
    
    # Make headers more readable
    readable_headers = [col.replace("_", " ").title() for col in columns]
    sheet.append(readable_headers)
    apply_header_style(sheet)
    
    # Data rows
    for row_data in data.get("data", []):
        row = [clean_text(row_data.get(col, "")) for col in columns]
        sheet.append(row)
    
    auto_adjust_column_width(sheet)
    return sheet


def create_sheet8(wb, data):
    """Create Sheet 8: Application Inventory Count."""
    sheet = wb.create_sheet("8-App Inventory")
    
    # Summary section
    summary = data.get("summary_counts", {})
    sheet.append(["Application Inventory Summary"])
    sheet.merge_cells("A1:B1")
    sheet["A1"].font = Font(bold=True, size=14)
    
    row = 2
    for key, value in summary.items():
        readable_key = key.replace("_", " ").title()
        sheet.append([readable_key, value])
        row += 1
    
    # Blank row
    sheet.append([])
    row += 1
    
    # Detailed inventory sections
    inventory = data.get("inventory", {})
    
    for section_name, items in inventory.items():
        if not items:
            continue
        
        # Section header
        section_row = sheet.max_row + 2
        readable_section = section_name.replace("_", " ").title()
        sheet.append([f"--- {readable_section} ({len(items)}) ---"])
        sheet.merge_cells(f"A{section_row}:D{section_row}")
        sheet[f"A{section_row}"].font = Font(bold=True, color="4472C4")
        
        # Get column headers from first item
        if items and isinstance(items[0], dict):
            headers = list(items[0].keys())
            readable_headers = [h.replace("_", " ").title() for h in headers]
            sheet.append(readable_headers)
            apply_header_style(sheet, sheet.max_row)
            
            for item in items:
                row_values = [clean_text(item.get(h, "")) for h in headers]
                sheet.append(row_values)
        
        # Blank row between sections
        sheet.append([])
    
    auto_adjust_column_width(sheet)
    return sheet


def create_sheet9_issues(wb, data):
    """Create Sheet 9A: Issues."""
    sheet = wb.create_sheet("9A-Issues")
    
    # Headers
    columns = data.get("issues_columns", [
        "filename", "class_name", "path", "issue_location",
        "type", "category", "description", "confidence", "mcm_node_id"
    ])
    readable_headers = [col.replace("_", " ").title() for col in columns]
    sheet.append(readable_headers)
    apply_header_style(sheet)
    
    # Data rows
    for issue in data.get("issues", []):
        row = [clean_text(issue.get(col, "")) for col in columns]
        sheet.append(row)
    
    # Merge cells for filename, class_name, path columns
    if len(data.get("issues", [])) > 0:
        merge_blocks(sheet, key_cols=(1, 2, 3), first_data_row=2)
    
    auto_adjust_column_width(sheet)
    return sheet


def create_sheet9_suggestions(wb, data):
    """Create Sheet 9B: Suggestions."""
    sheet = wb.create_sheet("9B-Suggestions")
    
    # Headers
    columns = data.get("suggestions_columns", [
        "filename", "class_name", "path", "suggestion", "mcm_node_id"
    ])
    readable_headers = [col.replace("_", " ").title() for col in columns]
    sheet.append(readable_headers)
    apply_header_style(sheet)
    
    # Data rows
    for suggestion in data.get("suggestions", []):
        row = [clean_text(suggestion.get(col, "")) for col in columns]
        sheet.append(row)
    
    # Merge cells for filename, class_name, path columns
    if len(data.get("suggestions", [])) > 0:
        merge_blocks(sheet, key_cols=(1, 2, 3), first_data_row=2)
    
    auto_adjust_column_width(sheet)
    return sheet


def create_summary_sheet(wb, sheet4_data, sheet7_data, sheet8_data, sheet9_data):
    """Create a summary/overview sheet."""
    sheet = wb.active
    sheet.title = "Overview"
    
    # Title
    sheet.append(["Technical Assessment Report"])
    sheet.merge_cells("A1:D1")
    sheet["A1"].font = Font(bold=True, size=16)
    sheet.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    sheet.append([])
    
    # Architecture Summary (from Sheet 4)
    if sheet4_data:
        summary = sheet4_data.get("summary", {})
        sheet.append(["Architecture Overview"])
        sheet[f"A{sheet.max_row}"].font = Font(bold=True, size=12)
        sheet.append(["Total Modules", summary.get("total_modules", 0)])
        sheet.append(["High Risk Modules", summary.get("high_risk", 0)])
        sheet.append(["Medium Risk Modules", summary.get("medium_risk", 0)])
        sheet.append(["Low Risk Modules", summary.get("low_risk", 0)])
        sheet.append(["Cyclic Dependencies", summary.get("cyclic_dependencies", 0)])
        sheet.append(["Boundary Violations", summary.get("boundary_violations", 0)])
        sheet.append([])
    
    # Code Quality Summary (from Sheet 7)
    if sheet7_data:
        summary = sheet7_data.get("summary", {})
        sheet.append(["Code Quality Summary"])
        sheet[f"A{sheet.max_row}"].font = Font(bold=True, size=12)
        sheet.append(["Total Files Analyzed", sheet7_data.get("total_files_analyzed", 0)])
        sheet.append(["High Risk Files", summary.get("high_risk_files", 0)])
        sheet.append(["Medium Risk Files", summary.get("medium_risk_files", 0)])
        sheet.append(["Low Risk Files", summary.get("low_risk_files", 0)])
        sheet.append(["Total Vulnerabilities", summary.get("total_vulnerabilities", 0)])
        sheet.append(["Total Bugs", summary.get("total_bugs", 0)])
        sheet.append(["Total Code Smells", summary.get("total_code_smells", 0)])
        sheet.append(["Total Lines of Code", summary.get("total_lines", 0)])
        sheet.append([])
    
    # Inventory Summary (from Sheet 8)
    if sheet8_data:
        counts = sheet8_data.get("summary_counts", {})
        sheet.append(["Application Inventory"])
        sheet[f"A{sheet.max_row}"].font = Font(bold=True, size=12)
        sheet.append(["Controllers", counts.get("controllers", 0)])
        sheet.append(["Services", counts.get("services", 0)])
        sheet.append(["Repositories", counts.get("repositories", 0)])
        sheet.append(["Classes", counts.get("classes", 0)])
        sheet.append(["Methods", counts.get("methods", 0)])
        sheet.append(["UI Screens", counts.get("ui_screens", 0)])
        sheet.append(["Third Party Dependencies", counts.get("third_party_dependencies", 0)])
        sheet.append([])
    
    # Issues Summary (from Sheet 9)
    if sheet9_data:
        summary = sheet9_data.get("summary", {})
        sheet.append(["Issues & Findings"])
        sheet[f"A{sheet.max_row}"].font = Font(bold=True, size=12)
        sheet.append(["Total Issues", summary.get("total_issues", 0)])
        sheet.append(["Total Suggestions", summary.get("total_suggestions", 0)])
        sheet.append(["Files with Issues", summary.get("total_files_with_issues", 0)])
        
        # Issues by type
        issues_by_type = summary.get("issues_by_type", {})
        for issue_type, count in issues_by_type.items():
            sheet.append([f"  - {issue_type}", count])
    
    auto_adjust_column_width(sheet)
    return sheet


def generate_excel():
    """Main function to generate the Excel report."""
    
    print("=" * 60)
    print("Excel Generator: Technical Assessment Report")
    print("=" * 60)
    
    # Load all JSON data
    print("\n📂 Loading sheet data...")
    
    sheet1_data = load_json_file(SHEET1_FILE)
    sheet2_data = load_json_file(SHEET2_FILE)
    sheet3_data = load_json_file(SHEET3_FILE)
    sheet4_data = load_json_file(SHEET4_FILE)
    sheet7_data = load_json_file(SHEET7_FILE)
    sheet8_data = load_json_file(SHEET8_FILE)
    sheet9_data = load_json_file(SHEET9_FILE)
    
    if not any([sheet1_data, sheet2_data, sheet3_data, sheet4_data, sheet7_data, sheet8_data, sheet9_data]):
        print("\n❌ No sheet data found. Run sheet generators first.")
        return False
    
    # Create workbook
    wb = Workbook()
    
    # Create Overview sheet (uses default active sheet)
    print("\n📝 Creating sheets...")
    create_summary_sheet(wb, sheet4_data, sheet7_data, sheet8_data, sheet9_data)
    print("  ✓ Overview sheet")
    
    SHEET_PIPELINE = [
        (sheet1_data, create_sheet1, "Sheet 1: Executive Summary", "data"),
        (sheet2_data, create_sheet2, "Sheet 2: Modernization", "data"),
        (sheet3_data, create_sheet3, "Sheet 3: App Summary", "data"),
        (sheet4_data, create_sheet4, "Sheet 4: Architecture", "data"),
        (sheet7_data, create_sheet7, "Sheet 7: Code Quality", "data"),
        (sheet8_data, create_sheet8, "Sheet 8: App Inventory", None),
        (sheet9_data, create_sheet9_issues, "Sheet 9A: Issues", "issues"),
        (sheet9_data, create_sheet9_suggestions, "Sheet 9B: Suggestions", "suggestions"),
    ]

    for data, creator, label, data_key in SHEET_PIPELINE:
        if data:
            creator(wb, data)
            count = len(data.get(data_key, [])) if data_key else ""
            print(f"  ✓ {label}" + (f" ({count} rows)" if count != "" else ""))
    
    # Ensure output directory exists
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    
    # Save workbook
    wb.save(OUTPUT_FILE)
    
    print("-" * 60)
    print(f"\n✅ Excel report generated: {OUTPUT_FILE}")
    print(f"   Sheets created: {len(wb.sheetnames)}")
    for name in wb.sheetnames:
        print(f"   - {name}")
    
    return True


if __name__ == "__main__":
    success = generate_excel()
    exit(0 if success else 1)

